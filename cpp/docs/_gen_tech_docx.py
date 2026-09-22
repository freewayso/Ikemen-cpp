# -*- coding: utf-8 -*-
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

out = r"e:\test\Ikemen-GO\cpp\docs\Ikemen-Cpp-Network-Tech.docx"


def set_run_font(run, size=11, bold=False, name_cn="微软雅黑", name_en="Calibri"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name_en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name_cn)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def add_heading_cn(doc, text, level):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 14, 3: 12}.get(level, 12), bold=True, name_cn="微软雅黑")
    return p


def add_p(doc, text, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=10, name_en="Consolas", name_cn="微软雅黑")
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    return p


def shade_header(cell):
    from docx.oxml import parse_xml

    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(
        r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="1F4E79"/>'
    )
    tcPr.append(shd)
    for p in cell.paragraphs:
        for run in p.runs:
            set_run_font(run, size=10, bold=True)
            run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        shade_header(cell)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.rows[r + 1].cells[c]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=9)
    doc.add_paragraph()
    return t


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.4)
        s.right_margin = Cm(2.4)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Ikemen C++ 运行时")
    set_run_font(r, size=22, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("网络大厅与帧同步转发技术文档")
    set_run_font(r, size=16, bold=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("版本：与 cpp/ 当前实现一致    文档位置：cpp/docs/")
    set_run_font(r, size=10)
    doc.add_paragraph()

    add_p(
        doc,
        "本文说明 C++ 客户端（ikemen_cpp）、登录大厅（ikemen_lobby）、对战转发（ikemen_relay）的架构、协议、存档、编译与联调。实现目录为 cpp/，大厅 schema 为 cpp/proto/lobby.proto。",
    )

    add_heading_cn(doc, "1 概述", 1)
    add_heading_cn(doc, "1.1 设计目标", 2)
    add_p(doc, "玩法侧采用 C++ 引擎加 Lua 角色逻辑，对齐 Ikemen GO 的战斗循环。联网拆成两段，互不混用同一套包：")
    add_p(doc, "大厅：TCP + Protocol Buffers 3。负责注册/登录、房间列表、创建与加入。每个房间最多 2 人。")
    add_p(doc, "对战：UDP + KCP 帧同步。relay 按房间隔离确认帧、历史缓冲与追帧。")
    add_p(
        doc,
        "大厅消息风格参考 luabus（命名消息、主循环驱动）。未接入 luna/lbus，序列化为自实现 proto3 线格式，不依赖 libprotobuf。",
    )
    add_heading_cn(doc, "1.2 不在本文范围", 2)
    add_p(doc, "完整 GGPO 回滚、ZSS 编译器、以及云厂商控制台以外的运维细则。")

    add_heading_cn(doc, "2 系统架构", 1)
    add_heading_cn(doc, "2.1 进程与端口", 2)
    add_table(
        doc,
        ["进程", "产物", "传输", "默认端口", "职责"],
        [
            ["ikemen_lobby", "bin_cpp/ikemen_lobby", "TCP", "8080", "登录、房间管理、START"],
            ["ikemen_relay", "bin_cpp/ikemen_relay", "UDP", "9000", "多房间 KCP 帧同步"],
            ["ikemen_cpp", "bin_cpp/ikemen_cpp", "TCP+UDP", "客户端", "登录 → 房间列表 → 对战"],
        ],
    )
    add_p(doc, "云主机需同时放行 TCP 8080 与 UDP 9000。")
    add_heading_cn(doc, "2.2 数据流", 2)
    add_code(
        doc,
        "客户端 A  ── TCP proto3 ──►  lobby  ── 房间列表 / START ──► 客户端 A / B\n"
        "客户端 A  ── UDP KCP IKFS ─►  relay  ◄── 客户端 B",
    )

    add_heading_cn(doc, "3 客户端流程", 1)
    add_p(doc, "普通启动（无 --host / --connect）按下列步骤：")
    steps = [
        "读取 data/net.ini，连接 Lobby 与 LobbyPort。",
        "进入 LOGIN：填写 USER（密码可空），点击 LOGIN 或 REGISTER。",
        "注册成功后自动登录，进入房间列表（S2C_ROOM_LIST 实时刷新）。",
        "按 C 创建房间，或选择 WAIT 房间按 Enter 加入；满 2 人后双方收到 START。",
        "用大厅房间 id（8 字符，如 r0000001）连接 relay，进入 KCP 对战。",
        "对战结束或 ESC：向大厅发送 LEAVE，回到房间列表。登录页 ESC 退出程序。",
    ]
    for i, s in enumerate(steps, 1):
        add_p(doc, "%d. %s" % (i, s))
    add_p(doc, "命令行可覆盖地址，避免写死 IP：")
    add_code(
        doc,
        "ikemen_cpp --lobby 127.0.0.1:8080 --relay 127.0.0.1:9000\n"
        "ikemen_cpp --host       # 跳过大厅，按 net.ini 的 Room 以 host 进 relay\n"
        "ikemen_cpp --connect    # 跳过大厅，以 guest 进 relay",
    )
    add_p(
        doc,
        "主菜单 NETWORK → HOST GAME / JOIN GAME 仍使用固定 Room=，仅适合本机快速测试。正式匹配走登录后的房间列表。",
    )

    add_heading_cn(doc, "4 配置文件", 1)
    add_p(doc, "仓库中 data/* 默认被 gitignore。本地或云上自行放置 data/net.ini：")
    add_code(
        doc,
        "[Net]\nRelay=127.0.0.1\nPort=9000\nRoom=kfm1\nLobby=127.0.0.1\nLobbyPort=8080",
    )
    add_p(doc, "Relay 可写成 ip:port。未单独配置 Lobby 时，大厅 IP 沿用 Relay 的地址。")

    add_heading_cn(doc, "5 大厅协议（TCP / proto3）", 1)
    add_p(doc, "定义：cpp/proto/lobby.proto。实现：lobby_proto.hpp、lobby.cpp、lobby_server.cpp。")
    add_heading_cn(doc, "5.1 帧格式", 2)
    add_code(doc, "[uint32 小端 载荷长度][Envelope 的 proto3 字节]")
    add_p(doc, "单帧载荷上限 65536 字节。Envelope.type 为消息号，Envelope.body 为对应 message 的序列化结果。")
    add_heading_cn(doc, "5.2 消息一览", 2)
    add_table(
        doc,
        ["type", "方向", "body", "说明"],
        [
            ["1 C2S_REGISTER", "C→S", "AuthReq", "注册，密码可空"],
            ["2 C2S_LOGIN", "C→S", "AuthReq", "登录"],
            ["3 C2S_LIST", "C→S", "空", "拉取房间列表"],
            ["4 C2S_CREATE", "C→S", "空", "创建房间，创建者 role=0"],
            ["5 C2S_JOIN", "C→S", "JoinReq", "加入 room_id"],
            ["6 C2S_LEAVE", "C→S", "空", "离开当前房间"],
            ["10 S2C_HELLO", "S→C", "Hello", "接入后立即发送"],
            ["11 S2C_ERROR", "S→C", "Error", "code + message"],
            ["12 S2C_LOGIN_OK", "S→C", "UserOk", "登录成功"],
            ["13 S2C_REGISTER_OK", "S→C", "UserOk", "随后客户端自动 LOGIN"],
            ["14 S2C_ROOM_LIST", "S→C", "RoomList", "全量快照"],
            ["15 S2C_JOIN_OK", "S→C", "JoinOk", "role：0 host / 1 guest"],
            ["16 S2C_START", "S→C", "JoinOk", "两人到齐，开始连 relay"],
            ["17 S2C_LEAVE_OK", "S→C", "空", "离开确认"],
        ],
    )
    add_heading_cn(doc, "5.3 房间信息 RoomInfo", 2)
    add_p(doc, "id：字母 r 加 7 位序号，共 8 字符，与 relay 房间键对齐。")
    add_p(doc, "status：0=WAIT 可加入；1=FIGHT 拒加入（ERR FIGHT）。")
    add_p(doc, "players / max_players：每房最多 2 人。host 为创建者用户名。")
    add_p(
        doc,
        "列表在登录成功、创建、加入、离开、掉线时向所有已登录客户端广播。未进房客户端约每 1.5 秒再发 C2S_LIST。",
    )
    add_heading_cn(doc, "5.4 错误码", 2)
    add_table(
        doc,
        ["code", "含义"],
        [
            ["BAD", "用户名非法"],
            ["EXISTS", "账号已注册"],
            ["AUTH", "登录失败"],
            ["ONLINE", "该账号已在线"],
            ["NEEDLOGIN", "未登录"],
            ["NOROOM", "房间不存在"],
            ["FULL", "房间已满"],
            ["FIGHT", "对战中不可加入"],
            ["MAX", "大厅房间数达上限"],
            ["CMD", "未知命令"],
        ],
    )
    add_heading_cn(doc, "5.5 限额", 2)
    add_p(doc, "用户名 1–16 位，字母数字下划线。密码 0–32 位，允许为空。大厅最多 64 个房间，每房 2 人。")

    add_heading_cn(doc, "6 账号存档", 1)
    add_p(doc, "路径默认 data/users.pb，不以明文保存密码。文件布局：")
    add_code(doc, '"IKUS" | uint32le(len) | UserStore（proto3） | sha256(UserStore)')
    add_p(doc, "Account 字段：user；16 字节 salt；pass_hash = sha256(salt 连接 pass)。校验失败则拒绝加载。")
    add_p(
        doc,
        "若不存在 users.pb，会尝试导入旧版 data/users.txt（每行 user 与可选密码）并写成 pb，之后只读写 pb。",
    )
    add_p(doc, "启动参数：ikemen_lobby [端口=8080] [存档路径=data/users.pb]")

    add_heading_cn(doc, "7 对战协议 IKFS（UDP / KCP）", 1)
    add_p(doc, "定义 relay_proto.hpp；服务 relay_server.cpp；客户端 room_sync.cpp。")
    add_heading_cn(doc, "7.1 多房间 KCP", 2)
    add_p(doc, "同一 UDP 端口上可同时存在多场对战。")
    add_p(
        doc,
        "conversation = 0x4B460000 | ((FNV-1a(room[8]) & 0x7FFF) << 1) | role。role 0 为 host，1 为 guest。",
    )
    add_p(doc, "会话按 UDP 源地址索引；房间状态按 8 字节 room 键隔离。")
    add_heading_cn(doc, "7.2 应用层头（32 字节，小端）", 2)
    add_table(
        doc,
        ["偏移", "长度", "字段"],
        [
            ["0", "4", "magic IKFS"],
            ["4", "1", "cmd"],
            ["5", "1", "role"],
            ["6", "8", "room"],
            ["14", "4", "frame"],
            ["18", "4", "a（INPUT=本端按键；CONFIRM=P0）"],
            ["22", "4", "b（CONFIRM=P1）"],
            ["26", "4", "seed"],
            ["30", "2", "count（CATCHUP_PACK 条数）"],
        ],
    )
    add_p(
        doc,
        "cmd：JOIN=10，JOIN_OK=11，INPUT=12，CONFIRM=13，CATCHUP=14，CATCHUP_PACK=15，DROP=16。",
    )
    add_p(
        doc,
        "CATCHUP_PACK 额外载荷每条 12 字节 (frame, i0, i1) 小端 u32，最多 40 条。Relay 约保留 3600 帧。空闲约 90 秒超时 DROP。同时对战场次上限 128。双方离开后删除该房间。",
    )
    add_p(doc, "仍兼容旧 IKRL（13 字节头）UDP 转发；正式对战使用 IKFS。")

    add_heading_cn(doc, "8 战斗同步", 1)
    add_p(doc, "双方每帧上传本端 input；relay 凑齐 P0 与 P1 后广播 CONFIRM。")
    add_p(doc, "丢帧或重连时客户端发送 CATCHUP，用历史包快进模拟到当前确认帧。")
    add_p(doc, "对战 HUD 显示 FPS、PING、帧号。角色逻辑位于 cpp/lua/chars/，引擎为 engine.cpp。")

    add_heading_cn(doc, "9 编译与部署", 1)
    add_p(doc, "Windows（MinGW，在 cpp/ 目录）：mingw32-make game relay lobby")
    add_p(doc, "Linux 默认只编服务端：make（relay + lobby）。CMake 见 cpp/CMakeLists.txt。")
    add_p(doc, "游戏依赖 SDL2、OpenGL、zlib、Lua 5.4.7、KCP。Linux 打包：bin_cpp/_pack_relay.sh。云上脚本：run.sh、run_lobby.sh。")

    add_heading_cn(doc, "10 本地联调", 1)
    add_code(
        doc,
        "bin_cpp/ikemen_lobby 8080 data/users.pb\n"
        "bin_cpp/ikemen_relay 9000\n"
        "bin_cpp/ikemen_cpp --lobby 127.0.0.1:8080 --relay 127.0.0.1:9000",
    )
    add_p(doc, "两个客户端使用不同用户名注册或登录；一方创建房间，另一方加入。")

    add_heading_cn(doc, "11 源码索引", 1)
    add_table(
        doc,
        ["路径", "内容"],
        [
            ["cpp/proto/lobby.proto", "大厅与存档 schema"],
            ["cpp/src/lobby_proto.hpp", "proto3 编解码"],
            ["cpp/src/lobby_server.cpp", "大厅服务"],
            ["cpp/src/lobby.cpp", "客户端 TCP"],
            ["cpp/src/sha256.hpp", "密码与文件校验"],
            ["cpp/src/relay_proto.hpp", "IKFS / KCP conv"],
            ["cpp/src/relay_server.cpp", "多房间 relay"],
            ["cpp/src/room_sync.cpp", "客户端帧同步"],
            ["cpp/src/ini.hpp", "net.ini"],
            ["cpp/src/engine.cpp", "画面流程与开战"],
            ["cpp/src/title.cpp", "登录 / 房间界面"],
        ],
    )

    add_heading_cn(doc, "12 已知限制", 1)
    add_p(doc, "登录界面为点阵英文字体，文案为英文。")
    add_p(doc, "密码使用 SHA-256 加盐，不是慢哈希（无 bcrypt / argon2），存档文件仍需操作系统权限保护。")
    add_p(doc, "data/users.pb、日志与二进制默认不纳入 git。")
    add_p(doc, "未编译官方 luabus 与 luna。若后续要在 Lua 侧接总线，应对齐同一套 Msg 与 TCP 帧格式。")

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
